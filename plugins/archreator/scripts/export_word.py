#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#     "python-docx>=1.1",
#     "pyyaml>=6.0",
# ]
# ///
"""Write one .docx from an explicit, ordered list of model documents.

A reader who will mark up prose in a tool everyone already has — Word, not
Markdown, not a graph — gets one document built from a named audience config:

    export_word.py --project . --config architecture/export/board.yml

An audience config is a small YAML file naming exactly what feeds the
document, in the order it should read:

    title: BigView — Board pack
    files:
      - architecture/README.md
      - architecture/1_estrategia/README.md
      - architecture/1_estrategia/*.md
      - architecture/2_negocio/*.md
    exclude:
      - architecture/2_negocio/actividades/*

**Nothing is included by default.** `files` is an explicit allow-list a human
wrote down, not a whole-model export with excludes — a glob expands in place,
sorted, but entries never reorder across each other, which is what lets a
project put a section out of its folder's own order without renaming files.
A path or glob that matches nothing is an error: an audience config that used
to work and now points at a renamed file fails loudly rather than shipping a
quiet hole.

**Every diagram travels, or the run fails.** Every fenced ```mermaid``` block
in a resolved source file is either rendered to an image (when `mmdc` —
mermaid-cli — happens to be on `PATH`) or embedded as a clearly labeled
source block. Either way it is counted, and the two counts must match before
the document is written — a diagram silently dropped by a parsing edge case
is exactly the failure this checks for. There is no bundled renderer: adding
one would be the headless-browser machinery a prior version of this method
built and then deleted for costing more to maintain than it was worth.

**The document is disposable**, same as every other generated artifact: it
lands under gitignored `.archreator/work/exports/`, stamped in its footer
with the project, the audience, when it was built and the model revision it
came from, so a copy that comes back edited can still be placed even if its
filename changed. Track Changes is on. Nothing here locks editing down
further — see `docs/adopting.md` § A Word document, for feedback.
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

# --------------------------------------------------------------------------
# This tool runs from the plugin and reads a project. The parse it needs -
# `model_graph.py` - lives in that project's `scripts/`, beside the two
# validators, because a project has to be able to check itself with no plugin
# installed and no network. Copied from `build_brief.py` rather than shared,
# exactly as `model.py` and `build_brief.py` already each carry their own copy.
# --------------------------------------------------------------------------
def _project_root(argv: list[str]) -> Path:
    for index, argument in enumerate(argv):
        if argument == "--project" and index + 1 < len(argv):
            return Path(argv[index + 1]).resolve()
        if argument.startswith("--project="):
            return Path(argument.split("=", 1)[1]).resolve()
    return Path.cwd().resolve()


def _find_parse(root: Path) -> Path | None:
    if not root.is_dir():
        return None
    for candidate in (root, *root.parents):
        parse = candidate / "scripts" / "model_graph.py"
        if parse.is_file():
            if candidate != root and not (root / "architecture").is_dir():
                return None
            return parse
        if (candidate / ".git").exists():
            return None
    return None


_ROOT = _project_root(sys.argv[1:])
_PARSE = _find_parse(_ROOT)
if _PARSE is None:
    if {"-h", "--help"} & set(sys.argv[1:]):
        print(__doc__)
        sys.exit(0)
    sys.exit(
        f"No archreator project at {_ROOT}: expected a directory holding "
        f"architecture/, with scripts/model_graph.py beside it or at the "
        f"enclosing repository root.\n"
        f"Run this from a project's root, or pass --project <path>."
    )
sys.path.insert(0, str(_PARSE.parent))

import model  # noqa: E402  (path is set up just above)

HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET = re.compile(r"^[-*]\s+(.*)$")
NUMBERED = re.compile(r"^\d+\.\s+(.*)$")
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
MERMAID_OPEN = re.compile(r"^```mermaid\s*$", re.MULTILINE)


def _is_glob(pattern: str) -> bool:
    return any(ch in pattern for ch in "*?[")


def load_config(path: Path) -> dict:
    import yaml

    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not data.get("files"):
        sys.exit(f"{path}: needs a non-empty 'files' list — the audience's "
                  f"explicit, ordered allow-list")
    data.setdefault("title", path.stem)
    data.setdefault("exclude", [])
    return data


def resolve_files(project: Path, config: dict, config_path: Path) -> list[Path]:
    """The ordered, deduplicated, excluded file list an audience declares.

    Order is `files`, top to bottom; a glob expands sorted in place but never
    reorders across entries — a project decides its own reading order rather
    than inheriting the filesystem's.
    """
    seen: set[Path] = set()
    ordered: list[Path] = []
    for entry in config["files"]:
        if _is_glob(entry):
            matches = sorted(p for p in project.glob(entry) if p.is_file())
            if not matches:
                sys.exit(f"{config_path}: '{entry}' matched no files under {project}")
        else:
            candidate = project / entry
            if not candidate.is_file():
                sys.exit(f"{config_path}: '{entry}' does not exist under {project}")
            matches = [candidate]
        for match in matches:
            if match not in seen:
                seen.add(match)
                ordered.append(match)

    excluded: set[Path] = set()
    for pattern in config["exclude"]:
        excluded.update(project.glob(pattern) if _is_glob(pattern) else [project / pattern])
    ordered = [f for f in ordered if f not in excluded]
    if not ordered:
        sys.exit(f"{config_path}: every matched file was excluded — nothing to export")
    return ordered


def count_mermaid(files: list[Path]) -> int:
    return sum(len(MERMAID_OPEN.findall(f.read_text(encoding="utf-8"))) for f in files)


def add_formatted(paragraph, text: str) -> None:
    """Bold, italic and inline code — the inline marks the model's prose uses."""
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def add_code_block(doc, code_text: str) -> None:
    from docx.shared import Pt

    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    run = para.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def render_mermaid(mermaid_text: str, index: int, image_dir: Path) -> Path | None:
    src = image_dir / f"diagram-{index}.mmd"
    dst = image_dir / f"diagram-{index}.png"
    src.write_text(mermaid_text, encoding="utf-8")
    try:
        result = subprocess.run(
            ["mmdc", "-i", str(src), "-o", str(dst), "-b", "white"],
            capture_output=True, text=True, timeout=60, check=False,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    return dst if result.returncode == 0 and dst.is_file() else None


def add_diagram(doc, mermaid_text: str, state: dict) -> None:
    from docx.shared import Inches

    state["diagrams_seen"] += 1
    image = (render_mermaid(mermaid_text, state["diagrams_seen"], state["image_dir"])
              if state["mmdc"] else None)
    if image:
        doc.add_picture(str(image), width=Inches(6))
    else:
        caption = doc.add_paragraph()
        caption.add_run(
            "Diagram source (Mermaid) — render at https://mermaid.live or in a "
            "Markdown viewer:"
        ).italic = True
        add_code_block(doc, mermaid_text)
    state["diagrams_embedded"] += 1


def add_table(doc, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:  # pragma: no cover - depends on the docx template
        pass
    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            add_formatted(cell.paragraphs[0], text)


def add_quote(doc, text: str) -> None:
    try:
        para = doc.add_paragraph(style="Quote")
    except KeyError:  # pragma: no cover - depends on the docx template
        para = doc.add_paragraph()
    add_formatted(para, text)
    for run in para.runs:
        run.italic = True


def is_table_row(line: str) -> bool:
    return "|" in line.strip()


def is_separator_row(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    cells = [c for c in cells if c]
    return bool(cells) and all(re.fullmatch(r":?-{1,}:?", c) for c in cells)


def split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def add_markdown(doc, text: str, state: dict) -> None:
    """The Markdown subset the model's own documents use — not CommonMark.

    Headings, paragraphs, bullet and numbered lists, pipe tables, blockquotes,
    fenced code (Mermaid handled specially) and inline bold/italic/code. What
    the layer documents and generated briefs actually contain, and nothing
    more — extending this to a general Markdown parser is not a cost this
    tool's job needs to carry.
    """
    lines = text.splitlines()
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            lang = stripped[3:].strip().lower()
            body = []
            i += 1
            while i < n and lines[i].strip() != "```":
                body.append(lines[i])
                i += 1
            i += 1  # past the closing fence
            code_text = "\n".join(body)
            if lang == "mermaid":
                add_diagram(doc, code_text, state)
            else:
                add_code_block(doc, code_text)
            continue

        if not stripped:
            i += 1
            continue

        heading = HEADING.match(stripped)
        if heading:
            level = len(heading.group(1))
            para = doc.add_heading("", level=min(level, 9))
            add_formatted(para, heading.group(2))
            i += 1
            continue

        if stripped.startswith(">"):
            add_quote(doc, stripped.lstrip(">").strip())
            i += 1
            continue

        bullet = BULLET.match(stripped)
        if bullet:
            para = doc.add_paragraph(style="List Bullet")
            add_formatted(para, bullet.group(1))
            i += 1
            continue

        numbered = NUMBERED.match(stripped)
        if numbered:
            para = doc.add_paragraph(style="List Number")
            add_formatted(para, numbered.group(1))
            i += 1
            continue

        if is_table_row(line) and i + 1 < n and is_separator_row(lines[i + 1]):
            header = split_row(line)
            i += 2
            rows = []
            while i < n and lines[i].strip() and is_table_row(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, header, rows)
            continue

        para = doc.add_paragraph()
        add_formatted(para, stripped)
        i += 1


def set_footer(doc, text: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    for section in doc.sections:
        paragraph = section.footer.paragraphs[0]
        run = paragraph.add_run(text)
        run.font.size = Pt(8)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


# `w:settings` fixes its children's order (ECMA-376 CT_Settings) — an element
# inserted out of sequence is exactly what makes Word offer to "repair" a
# document, which is the last thing a document meant to earn a stakeholder's
# trust should do. This is python-docx's own internal ordering
# (`docx.oxml.settings.CT_Settings._tag_seq`, not exposed at runtime, so
# copied here) — enough of it to place `w:trackRevisions` correctly relative
# to whatever a freshly created `Document()` already has.
_SETTINGS_ORDER = (
    "w:writeProtection", "w:view", "w:zoom", "w:removePersonalInformation",
    "w:removeDateAndTime", "w:doNotDisplayPageBoundaries", "w:displayBackgroundShape",
    "w:printPostScriptOverText", "w:printFractionalCharacterWidth", "w:printFormsData",
    "w:embedTrueTypeFonts", "w:embedSystemFonts", "w:saveSubsetFonts", "w:saveFormsData",
    "w:mirrorMargins", "w:alignBordersAndEdges", "w:bordersDoNotSurroundHeader",
    "w:bordersDoNotSurroundFooter", "w:gutterAtTop", "w:hideSpellingErrors",
    "w:hideGrammaticalErrors", "w:activeWritingStyle", "w:proofState", "w:formsDesign",
    "w:attachedTemplate", "w:linkStyles", "w:stylePaneFormatFilter", "w:stylePaneSortMethod",
    "w:documentType", "w:mailMerge", "w:revisionView", "w:trackRevisions",
    "w:doNotTrackMoves", "w:doNotTrackFormatting", "w:documentProtection",
    "w:autoFormatOverride", "w:styleLockTheme", "w:styleLockQFSet", "w:defaultTabStop",
    "w:autoHyphenation", "w:consecutiveHyphenLimit", "w:hyphenationZone",
    "w:doNotHyphenateCaps", "w:showEnvelope", "w:summaryLength", "w:clickAndTypeStyle",
    "w:defaultTableStyle", "w:evenAndOddHeaders", "w:bookFoldRevPrinting",
    "w:bookFoldPrinting", "w:bookFoldPrintingSheets", "w:drawingGridHorizontalSpacing",
    "w:drawingGridVerticalSpacing", "w:displayHorizontalDrawingGridEvery",
    "w:displayVerticalDrawingGridEvery", "w:doNotUseMarginsForDrawingGridOrigin",
    "w:drawingGridHorizontalOrigin", "w:drawingGridVerticalOrigin", "w:doNotShadeFormData",
    "w:noPunctuationKerning", "w:characterSpacingControl", "w:printTwoOnOne",
    "w:strictFirstAndLastChars", "w:noLineBreaksAfter", "w:noLineBreaksBefore",
    "w:savePreviewPicture", "w:doNotValidateAgainstSchema", "w:saveInvalidXml",
    "w:ignoreMixedContent", "w:alwaysShowPlaceholderText", "w:doNotDemarcateInvalidXml",
    "w:saveXmlDataOnly", "w:useXSLTWhenSaving", "w:saveThroughXslt", "w:showXMLTags",
    "w:alwaysMergeEmptyNamespace", "w:updateFields", "w:hdrShapeDefaults", "w:footnotePr",
    "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr", "w:attachedSchema",
    "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
    "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions", "w:readModeInkLockDown",
    "w:smartTagType", "sl:schemaLibrary", "w:shapeDefaults", "w:doNotEmbedSmartTags",
    "w:decimalSymbol", "w:listSeparator",
)


def enable_track_changes(doc) -> None:
    """Turn Track Changes on.

    `w:trackRevisions` is the actual element name — "Track Changes" is only
    the feature's name in Word's UI. Inserted right before the first existing
    child that the schema sequence says comes after it, or appended at the
    end if nothing does; never at a fixed index, which would misplace it the
    moment the python-docx template changes what it ships by default.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = doc.settings.element
    track = OxmlElement("w:trackRevisions")
    later = {qn(tag) for tag in _SETTINGS_ORDER[_SETTINGS_ORDER.index("w:trackRevisions") + 1:]}
    successor = next((child for child in settings if child.tag in later), None)
    if successor is not None:
        successor.addprevious(track)
    else:
        settings.append(track)


def build_document(project: Path, config: dict, files: list[Path]) -> tuple[object, dict]:
    from docx import Document

    doc = Document()
    doc.add_heading(config["title"], level=0)

    state = {
        "diagrams_seen": 0,
        "diagrams_embedded": 0,
        "mmdc": shutil.which("mmdc") is not None,
        "image_dir": Path(tempfile.mkdtemp(prefix="archreator-export-")),
    }
    for source in files:
        caption = doc.add_paragraph()
        run = caption.add_run(str(source.relative_to(project)))
        run.italic = True
        run.font.size = doc.styles["Normal"].font.size
        add_markdown(doc, source.read_text(encoding="utf-8"), state)
    return doc, state


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--project", default=".",
                        help="the project to read (default: the working directory)")
    parser.add_argument("--config", required=True, type=Path,
                        help="the audience config naming what feeds the document")
    parser.add_argument("--to", type=Path,
                        help="where to write the .docx (default: .archreator/work/exports/)")
    args = parser.parse_args()

    project = Path(args.project).resolve()
    config_path = args.config.resolve()
    if not config_path.is_file():
        sys.exit(f"No config at {config_path}")

    config = load_config(config_path)
    files = resolve_files(project, config, config_path)
    diagram_total = count_mermaid(files)

    doc, state = build_document(project, config, files)

    if state["diagrams_embedded"] != diagram_total:
        sys.exit(
            f"Diagram count mismatch: {diagram_total} mermaid block(s) in the source "
            f"files, {state['diagrams_embedded']} made it into the document. Refusing "
            f"to write a document that silently dropped a diagram."
        )

    when = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    rev = model.revision()
    set_footer(
        doc,
        f"{project.name} · {config['title']} · generated {when}"
        + (f" · revision {rev[:12]}" if rev else ""),
    )
    enable_track_changes(doc)

    out_dir = args.to or (project / ".archreator" / "work" / "exports")
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    out_path = out_dir / f"{config_path.stem}-{stamp}.docx"
    doc.save(str(out_path))

    rendered = state["mmdc"]
    print(
        f"Word export written to {out_path} — {len(files)} file(s), "
        f"{diagram_total} diagram(s) ({'rendered' if rendered else 'kept as source, no mmdc on PATH'}), "
        f"Track Changes on."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
